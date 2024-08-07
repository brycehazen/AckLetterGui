import pandas as pd
from docx import Document
import os
import glob
from docxcompose.composer import Composer
from docx import Document as Document_compose
from PySide6.QtCore import QObject, Signal, QThread, QIODevice 
from PySide6.QtWidgets import QTextEdit
from PySide6.QtGui import QTextCursor

class Logger(QObject):
    log_signal = Signal(str, bool)

    def __init__(self):
        super().__init__()

    def log(self, message, update_only=False):
        self.log_signal.emit(message, update_only)

class Worker(QThread):
    progress = Signal(str, bool)
    finished = Signal(str)

    def __init__(self, output_dir, template_path):
        super().__init__()
        self.output_dir = output_dir
        self.template_path = template_path

    def run(self):
        try:
            latest_csv = find_latest_complete_csv(self.output_dir)
            self.progress.emit(f"Latest CSV file found: {latest_csv}", False)
            logger = Logger()
            logger.log_signal.connect(self.log)
            mail_merge = MailMerge(self.output_dir, logger)
            mail_merge.merge(latest_csv, self.template_path)
            self.finished.emit("")
        except Exception as e:
            self.finished.emit(f"An error occurred during mail merge: {e}")

    def log(self, message, update_only=False):
        self.progress.emit(message, update_only)

class MailMerge:
    def __init__(self, output_dir, logger):
        self.output_dir = output_dir
        self.logger = logger

    def read_csv_file(self, file_path):
        """Attempt to read a CSV file with multiple encodings."""
        try:
            # Try reading with UTF-8 first
            return pd.read_csv(file_path, encoding='utf-8'), 'utf-8'
        except UnicodeDecodeError:
            # If UTF-8 fails, try ISO-8859-1
            return pd.read_csv(file_path, encoding='ISO-8859-1'), 'ISO-8859-1'

    def merge(self, data_path, template_path):
        # Read the data
        df, encoding = self.read_csv_file(data_path)

        self.logger.log(f"Starting mail merge process")
        self.logger.log(f"{len(df)} mail merges will be performed")
        self.logger.log(f"Beginning mail merge using '{os.path.basename(template_path)}'")

        # Process each row in the data
        for index, row in df.iterrows():
            # Load the template document
            doc = Document(template_path)

            # Replace placeholders with actual data
            for paragraph in doc.paragraphs:
                self._replace_placeholders(paragraph, row)

            # Save the personalized document
            output_path = os.path.join(self.output_dir, f"merged_letter_{index}.docx")
            doc.save(output_path)

            # Update progress
            self.logger.log(f"Mail merged completed {index + 1} out of {len(df)}", update_only=True)

        self.logger.log("Cleaning up mail merge files")

        # Combine individual documents into a single document
        self.combine_documents(template_path)

        # Clean up individual merged files after combining
        self.cleanup_individual_files()

        self.logger.log("Mail merge complete")

    def _replace_placeholders(self, paragraph, data):
        for key, value in data.items():
            if key in paragraph.text:
                paragraph.text = paragraph.text.replace(f"«{key}»", str(value))

    def combine_documents(self, template_path):
        files = sorted(glob.glob(os.path.join(self.output_dir, 'merged_letter_*.docx')))
        merged_document = Document_compose(files[0])
        composer = Composer(merged_document)

        for file in files[1:]:
            doc = Document_compose(file)
            composer.append(doc)

        base_name = os.path.basename(template_path)
        output_path = os.path.join(self.output_dir, f"Merged_{base_name}")
        composer.save(output_path)

    def cleanup_individual_files(self):
        files = glob.glob(os.path.join(self.output_dir, 'merged_letter_*.docx'))
        for file in files:
            os.remove(file)


def find_docx_template(input_dir):
    list_of_files = glob.glob(os.path.join(input_dir, '*.docx'))
    if len(list_of_files) == 1:
        return list_of_files[0]
    else:
        raise FileNotFoundError("Exactly one .docx template must be present in the directory.")

def find_latest_complete_csv(output_dir):
    list_of_files = glob.glob(os.path.join(output_dir, '*_complete.csv'))  # Adjust the pattern if necessary
    latest_file = max(list_of_files, key=os.path.getctime)
    return latest_file
